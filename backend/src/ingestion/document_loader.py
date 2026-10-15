# 将全部文件加载为Document对象

from pathlib import Path
from typing import List
from langchain_core.documents import Document
from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    TextLoader,
)

from pptx import Presentation


# =========================
# 支持的文件类型
# =========================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
}


# =========================
# PDF Loader
# =========================

def load_pdf(file_path: Path) -> List[Document]:

    loader = PyPDFLoader(str(file_path))

    documents = loader.load()

    return documents


# =========================
# Word Loader
# =========================

def load_docx(file_path: Path) -> List[Document]:

    loader = Docx2txtLoader(str(file_path))

    documents = loader.load()

    return documents


# =========================
# PPT Loader
# 一张幻灯片 = 一个 Document
# =========================

def load_pptx(file_path: Path) -> List[Document]:
    presentation = Presentation(str(file_path))
    documents = []
    for slide_index, slide in enumerate(presentation.slides):
        texts = []
        for shape in slide.shapes:

            # 普通文本框
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)

            # PPT中的表格
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text.strip())
                    texts.append(" | ".join(row_text))

        page_content = "\n".join(texts)
        if not page_content.strip():
            continue
        document = Document(
            page_content=page_content,
            metadata={
                "source": str(file_path),
                "file_name": file_path.name,
                "file_type": "pptx",
                "slide": slide_index + 1,
            }
        )
        documents.append(document)
    return documents

# =========================
# TXT / Markdown
# =========================

def load_text(file_path: Path) -> List[Document]:
    loader = TextLoader(
        str(file_path),
        encoding="utf-8"
    )
    return loader.load()

# =========================
# 单文件统一入口
# =========================

def load_single_file(file_path: Path) -> List[Document]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        documents = load_pdf(file_path)

    elif suffix == ".docx":
        documents = load_docx(file_path)

    elif suffix == ".pptx":
        documents = load_pptx(file_path)

    elif suffix in {".txt", ".md"}:
        documents = load_text(file_path)

    else:
        print(f"[跳过] 暂不支持文件：{file_path}")
        return []

    # =========================
    # 给所有Document补充统一Metadata
    # =========================

    for document in documents:
        document.metadata["file_name"] = file_path.name
        document.metadata["file_type"] = suffix.replace(".", "")
        document.metadata["source"] = str(file_path)

    return documents

# =========================
# 加载整个企业知识库目录
# =========================

def load_knowledge_base(
        knowledge_base_path: str
) -> List[Document]:
    root_path = Path(knowledge_base_path)
    all_documents = []

    if not root_path.exists():
        raise FileNotFoundError(
            f"知识库目录不存在：{root_path}"
        )

    for file_path in root_path.rglob("*"):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        print(
            f"\n正在加载：{file_path}"
        )
        try:

            documents = load_single_file(
                file_path
            )

            # =========================
            # 根据目录自动增加业务Domain
            # =========================

            relative_path = file_path.relative_to(
                root_path
            )

            if len(relative_path.parts) > 1:
                domain = relative_path.parts[0]

            else:
                domain = "general"

            for document in documents:
                document.metadata["domain"] = domain

            all_documents.extend(
                documents
            )

            print(
                f"加载成功：{len(documents)} 个 Document"
            )

        except Exception as e:

            print(
                f"[加载失败] {file_path}"
            )

            print(
                f"错误原因：{e}"
            )

    return all_documents